"""
S5: Eingaben der Verlustmessung erzeugen (deterministisch).

    dienst/.venv/bin/python scripts/s5-quellen.py docs/belege/S5/eingabe

Schreibt in das Zielverzeichnis:
  bilder/dachstuhl.png, bilder/feuchte-diagramm.png, bilder/zierlinie.png   Bilder (Pillow, fester Inhalt)
  referenz.fodt     Prüfbericht als Flat-XML-ODF-Textdokument mit Standardvorlagen einer Textverarbeitung
                    (Titel, Überschrift 1/2, Textkörper, Listen, Tabellen mit Kopfzeile und Spannen,
                    Beschriftungen mit text:sequence, drei Bilder, drei Fußnoten, Textmarken, Querverweise)
  referenz.md       derselbe Bericht in Markdown (Vorlage „bericht“); Markdown kennt keine verbundenen Zellen
  erwartet.json     Grundwahrheit: Abschnitte, Tabellen (Zellen, Spannen, Beschriftung), Bilder
                    (Alternativtext, Beschriftung, dekorativ), Fußnoten, Listen, Verweise, Fließtext
  grenzfall.docx    Word-Dokument aus der Normalvorlage von Word (python-docx) mit Inhalten, die das
                    Mindestprofil nicht abbilden kann: Kommentar, Änderungsverfolgung, Textfeld, Formel,
                    Kopf-/Fußzeile, Inhaltsverzeichnisfeld, Seitenumbruch, BMP-Bild, Bild ohne
                    Alternativtext, verschachtelte Tabelle, Inhaltssteuerelement, mailto-Link,
                    REF/PAGEREF-Felder, Fußnote, getippte Spiegelstriche, direkt formatierte „Überschrift“
  grenzfall-erwartet.json   was mit jedem dieser Merkmale geschehen soll

referenz.docx und referenz.pdf erzeugt danach LibreOffice aus referenz.fodt (scripts/belege-s5.sh).
"""
import base64
import copy
import io
import json
import os
import sys
import zipfile

from PIL import Image, ImageDraw, ImageFont

OUT = sys.argv[1] if len(sys.argv) > 1 else 'docs/belege/S5/eingabe'
FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
ID_REFERENZ = 'urn:uuid:7d1e5c3a-4b2f-4e6a-9c8d-0f1a2b3c4d5e'
ID_GRENZFALL = 'urn:uuid:9a8b7c6d-5e4f-4a3b-8c2d-1e0f9a8b7c6d'
DATUM = '2026-10-13T08:00:00Z'

TITEL = 'Prüfbericht Dachstuhl Haus Nr. 12'
ALT_FOTO = 'Foto des Dachstuhls von innen mit dunklen Feuchteflecken am Sparren Nord'
CAP_FOTO = 'Abbildung 1: Dachstuhl, Blick nach Norden'
ALT_DIAGRAMM = 'Balkendiagramm der Holzfeuchte: M1 24 %, M2 14 %, M3 19 %, M4 27 %'
CAP_DIAGRAMM = 'Abbildung 2: Holzfeuchte je Messpunkt'
CAP_T1 = 'Tabelle 1: Holzfeuchte je Messpunkt'
CAP_T2 = 'Tabelle 2: Kosten der Sanierung in Euro'
FN = [
    'Messung mit Widerstandsmessgerät nach ÖNORM B 3802-2.',
    'Grenzwert für tragendes Bauholz: 20 Prozent Holzfeuchte.',
    'Abstützung nach Angabe des Statikers vom 10. Oktober 2026.',
]
T1_KOPF = ['Messpunkt', 'Bauteil', 'Holzfeuchte %', 'Bewertung']
T1 = [['M1', 'Sparren Nord', '24', 'zu feucht'], ['M2', 'Sparren Süd', '14', 'in Ordnung'], ['M3', 'Pfette', '19', 'Grenzwert'], ['M4', 'Mauerbank', '27', 'zu feucht']]
T2_KOPF = ['Gewerk', 'Leistung', 'Betrag']
# Logische Zellen: (zeile, spalte, text, colspan, rowspan); Zeile 0 = Kopf.
T2_ZELLEN = [
    (1, 0, 'Zimmerei', 1, 1), (1, 1, 'Sparren ersetzen', 1, 1), (1, 2, '4.800,00', 1, 1),
    (2, 0, 'Spengler', 1, 2), (2, 1, 'Rinne erneuern', 1, 1), (2, 2, '1.250,00', 1, 1),
    (3, 1, 'Anschluss Kamin', 1, 1), (3, 2, '640,00', 1, 1),
    (4, 0, 'Summe netto', 2, 1), (4, 2, '6.690,00', 1, 1),
]
P_ZUSAMMENFASSUNG = 'Der Dachstuhl ist tragfähig. An der Nordseite sind Sparren und Mauerbank zu feucht; die Sanierung kostet rund 6.700 Euro netto.'
P_HOLZ = 'Kein Befall durch Hausbock oder Anobien festgestellt.'
P_EMPFEHLUNG = 'Die feuchten Bauteile sind zu trocknen oder zu ersetzen. Normen und Merkblätter: '
URL = 'https://www.austrian-standards.at/'
LISTE_FEUCHTE = [('Sparren Nord', ['Flecken auf 1,2 m Länge', 'Oberfläche weich']), ('Mauerbank Nord', []), ('Pfette, nur im Randbereich', [])]
LISTE_MASSNAHMEN = ['Rinne an der Nordseite abdichten', 'Sparren Nord abstützen', 'Raum täglich lüften']
SCHLUSS = 'Innsbruck, 13. Oktober 2026'


def font(size):
    return ImageFont.truetype(FONT, size)


def png(img):
    b = io.BytesIO()
    img.save(b, 'PNG', optimize=False)
    return b.getvalue()


def bilder():
    # Foto-ähnliches Bild: Dachstuhl mit Sparren und dunklen Flecken.
    img = Image.new('RGB', (480, 300), (210, 190, 160))
    d = ImageDraw.Draw(img)
    for y in range(300):
        d.line([(0, y), (480, y)], fill=(200 - y // 6, 180 - y // 6, 150 - y // 8))
    for x in range(20, 480, 70):
        d.polygon([(x, 300), (x + 18, 300), (240 + (x - 240) // 6 + 9, 20), (240 + (x - 240) // 6, 20)], fill=(150, 110, 70))
    d.rectangle([0, 270, 480, 300], fill=(120, 90, 60))
    for cx, cy, r in [(60, 200, 22), (75, 160, 15), (40, 280, 18), (300, 285, 10)]:
        d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=(70, 55, 40))
    foto = png(img)
    # Balkendiagramm mit Beschriftung.
    img = Image.new('RGB', (480, 300), 'white')
    d = ImageDraw.Draw(img)
    d.line([(50, 260), (460, 260)], fill='black', width=2)
    d.line([(50, 20), (50, 260)], fill='black', width=2)
    gy = 260 - 20 * 8
    d.line([(50, gy), (460, gy)], fill=(200, 0, 0), width=1)
    d.text((400, gy - 16), '20 %', fill=(200, 0, 0), font=font(12))
    for i, (name, wert) in enumerate([('M1', 24), ('M2', 14), ('M3', 19), ('M4', 27)]):
        x = 80 + i * 95
        d.rectangle([x, 260 - wert * 8, x + 50, 260], fill=(60, 100, 160))
        d.text((x + 12, 265), name, fill='black', font=font(14))
        d.text((x + 10, 260 - wert * 8 - 18), f'{wert} %', fill='black', font=font(13))
    diagramm = png(img)
    img = Image.new('RGB', (600, 8), 'white')
    ImageDraw.Draw(img).rectangle([0, 3, 600, 4], fill=(120, 120, 120))
    zier = png(img)
    return foto, diagramm, zier


def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')


def fodt(foto, diagramm, zier):
    b64 = lambda data: base64.b64encode(data).decode('ascii')
    note = lambda i: (f'<text:note text:id="ftn{i + 1}" text:note-class="footnote"><text:note-citation>{i + 1}</text:note-citation>'
                      f'<text:note-body><text:p text:style-name="Footnote">{esc(FN[i])}</text:p></text:note-body></text:note>')

    def cell(t, style='Table_20_Contents', extra=''):
        return f'<table:table-cell table:style-name="Zelle" office:value-type="string"{extra}><text:p text:style-name="{style}">{esc(t)}</text:p></table:table-cell>'

    def bild(name, data, w, h, alt, caption_no, caption):
        # Wie in Word: Bildabsatz (Bild als Zeichen) und darunter der Beschriftungsabsatz mit Nummernkreis.
        return (f'<text:p text:style-name="Text_20_body"><draw:frame draw:style-name="fr2" draw:name="{name}" text:anchor-type="as-char" svg:width="{w}cm" svg:height="{h}cm" draw:z-index="{caption_no}">'
                f'<draw:image draw:mime-type="image/png"><office:binary-data>{b64(data)}</office:binary-data></draw:image><svg:title>{esc(alt)}</svg:title><svg:desc>{esc(alt)}</svg:desc></draw:frame></text:p>'
                f'<text:p text:style-name="Illustration">Abbildung <text:sequence text:ref-name="refIllustration{caption_no - 1}" text:name="Abbildung" text:formula="ooow:Abbildung+1" style:num-format="1">{caption_no}</text:sequence>: {esc(caption.split(": ", 1)[1])}</text:p>')

    t1 = ''.join(f'<table:table-row>{"".join(cell(c) for c in r)}</table:table-row>' for r in T1)
    t2_rows = []
    for zeile in range(1, 5):
        cells = ''
        for spalte in range(3):
            z = [c for c in T2_ZELLEN if c[0] == zeile and c[1] == spalte]
            if z:
                _, _, t, cs, rs = z[0]
                extra = (f' table:number-columns-spanned="{cs}"' if cs > 1 else '') + (f' table:number-rows-spanned="{rs}"' if rs > 1 else '')
                cells += cell(t, extra=extra)
            else:
                cells += '<table:covered-table-cell/>'
        t2_rows.append(f'<table:table-row>{cells}</table:table-row>')

    def liste(items, style, nested_style=None):
        out = f'<text:list text:style-name="{style}">'
        for t, sub in items:
            out += f'<text:list-item><text:p text:style-name="List_20_Contents">{t}</text:p>'
            if sub:
                out += '<text:list>' + ''.join(f'<text:list-item><text:p text:style-name="List_20_Contents">{esc(s)}</text:p></text:list-item>' for s in sub) + '</text:list>'
            out += '</text:list-item>'
        return out + '</text:list>'

    body = f'''
      <text:p text:style-name="Title">{esc(TITEL)}</text:p>
      <text:h text:style-name="Heading_20_1" text:outline-level="1">Zusammenfassung</text:h>
      <text:p text:style-name="Text_20_body">{esc(P_ZUSAMMENFASSUNG)}</text:p>
      <text:h text:style-name="Heading_20_1" text:outline-level="1"><text:bookmark-start text:name="Befund"/>Befund<text:bookmark-end text:name="Befund"/></text:h>
      <text:p text:style-name="Text_20_body">Begehung am 9. Oktober 2026 mit dem Eigentümer. Die Holzfeuchte wurde an vier Punkten gemessen{note(0)}; die Werte stehen im Abschnitt <text:bookmark-ref text:reference-format="text" text:ref-name="Messwerte">Messwerte</text:bookmark-ref> und in <text:sequence-ref text:reference-format="category-and-value" text:ref-name="refTable0">Tabelle 1</text:sequence-ref>.</text:p>
      <text:h text:style-name="Heading_20_2" text:outline-level="2">Feuchteschäden</text:h>
      <text:p text:style-name="Text_20_body">Feuchteflecken am Sparren Nord und an der Mauerbank, siehe <text:sequence-ref text:reference-format="category-and-value" text:ref-name="refIllustration0">Abbildung 1</text:sequence-ref>.</text:p>
      {bild('Foto', foto, 12, 7.5, ALT_FOTO, 1, CAP_FOTO)}
      {liste([(esc(t), s) for t, s in LISTE_FEUCHTE], 'List_20_1')}
      <text:h text:style-name="Heading_20_2" text:outline-level="2">Holzschädlinge</text:h>
      <text:p text:style-name="Text_20_body">{esc(P_HOLZ)}</text:p>
      <text:h text:style-name="Heading_20_1" text:outline-level="1"><text:bookmark-start text:name="Messwerte"/>Messwerte<text:bookmark-end text:name="Messwerte"/></text:h>
      <text:p text:style-name="Text_20_body">Gemessen in 2 cm Tiefe bei 8 °C Lufttemperatur{note(1)}.</text:p>
      <table:table table:name="Messwerte_Tabelle" table:style-name="Tab">
        <table:table-column table:number-columns-repeated="4"/>
        <table:table-header-rows><table:table-row>{"".join(cell(c, 'Table_20_Heading') for c in T1_KOPF)}</table:table-row></table:table-header-rows>
        {t1}
      </table:table>
      <text:p text:style-name="Table">Tabelle <text:sequence text:ref-name="refTable0" text:name="Tabelle" text:formula="ooow:Tabelle+1" style:num-format="1">1</text:sequence>: {esc(CAP_T1.split(": ", 1)[1])}</text:p>
      {bild('Diagramm', diagramm, 12, 7.5, ALT_DIAGRAMM, 2, CAP_DIAGRAMM)}
      <text:h text:style-name="Heading_20_1" text:outline-level="1">Kosten</text:h>
      <table:table table:name="Kosten" table:style-name="Tab">
        <table:table-column table:number-columns-repeated="3"/>
        <table:table-header-rows><table:table-row>{"".join(cell(c, 'Table_20_Heading') for c in T2_KOPF)}</table:table-row></table:table-header-rows>
        {"".join(t2_rows)}
      </table:table>
      <text:p text:style-name="Table">Tabelle <text:sequence text:ref-name="refTable1" text:name="Tabelle" text:formula="ooow:Tabelle+1" style:num-format="1">2</text:sequence>: {esc(CAP_T2.split(": ", 1)[1])}</text:p>
      <text:h text:style-name="Heading_20_1" text:outline-level="1">Empfehlung</text:h>
      <text:p text:style-name="Text_20_body">{esc(P_EMPFEHLUNG)}<text:a xlink:type="simple" xlink:href="{URL}" text:style-name="Internet_20_link">Austrian Standards</text:a>.</text:p>
      <text:h text:style-name="Heading_20_2" text:outline-level="2">Sofortmaßnahmen</text:h>
      <text:list text:style-name="Numbering_20_123">
        <text:list-item><text:p text:style-name="List_20_Contents">{esc(LISTE_MASSNAHMEN[0])}</text:p></text:list-item>
        <text:list-item><text:p text:style-name="List_20_Contents">{esc(LISTE_MASSNAHMEN[1])}{note(2)}</text:p></text:list-item>
        <text:list-item><text:p text:style-name="List_20_Contents">{esc(LISTE_MASSNAHMEN[2])}</text:p></text:list-item>
      </text:list>
      <text:p text:style-name="Text_20_body"><draw:frame draw:style-name="fr2" draw:name="Zierlinie" draw:decorative="true" text:anchor-type="as-char" svg:width="15cm" svg:height="0.2cm" draw:z-index="2"><draw:image draw:mime-type="image/png"><office:binary-data>{b64(zier)}</office:binary-data></draw:image></draw:frame></text:p>
      <text:p text:style-name="Text_20_body">{esc(SCHLUSS)}</text:p>'''
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
  xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
  xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
  xmlns:xlink="http://www.w3.org/1999/xlink"
  xmlns:dc="http://purl.org/dc/elements/1.1/"
  xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0"
  xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
  office:version="1.3" office:mimetype="application/vnd.oasis.opendocument.text">
  <!-- S5-Referenz: erzeugt von scripts/s5-quellen.py. Struktur nur aus Vorlagen, Gliederungsebenen,
       Listen- und Tabellendefinitionen (Teil 0, Abschnitt 15.2). -->
  <office:meta>
    <dc:title>{esc(TITEL)}</dc:title>
    <dc:language>de-AT</dc:language>
    <dc:date>{DATUM}</dc:date>
    <dc:creator>Büro Beispiel Bautechnik</dc:creator>
    <meta:user-defined meta:name="hpro:id">{ID_REFERENZ}</meta:user-defined>
  </office:meta>
  <office:font-face-decls>
    <style:font-face style:name="DejaVu Serif" svg:font-family="'DejaVu Serif'" style:font-family-generic="roman"/>
    <style:font-face style:name="DejaVu Sans" svg:font-family="'DejaVu Sans'" style:font-family-generic="swiss"/>
  </office:font-face-decls>
  <office:styles>
    <style:default-style style:family="paragraph"><style:text-properties style:font-name="DejaVu Serif" fo:font-size="11pt" fo:language="de" fo:country="AT"/></style:default-style>
    <style:style style:name="Standard" style:family="paragraph" style:class="text"/>
    <style:style style:name="Heading" style:family="paragraph" style:parent-style-name="Standard" style:class="text"><style:text-properties style:font-name="DejaVu Sans" fo:font-weight="bold"/></style:style>
    <style:style style:name="Text_20_body" style:display-name="Text body" style:family="paragraph" style:parent-style-name="Standard" style:class="text"><style:paragraph-properties fo:margin-bottom="0.25cm"/></style:style>
    <style:style style:name="List_20_Contents" style:display-name="List Contents" style:family="paragraph" style:parent-style-name="Standard" style:class="html"/>
    <style:style style:name="Caption" style:family="paragraph" style:parent-style-name="Standard" style:class="extra"><style:text-properties fo:font-style="italic"/></style:style>
    <style:style style:name="Illustration" style:family="paragraph" style:parent-style-name="Caption" style:class="extra"/>
    <style:style style:name="Table" style:family="paragraph" style:parent-style-name="Caption" style:class="extra"/>
    <style:style style:name="Title" style:family="paragraph" style:parent-style-name="Heading" style:class="chapter"><style:text-properties fo:font-size="20pt"/></style:style>
    <style:style style:name="Heading_20_1" style:display-name="Heading 1" style:family="paragraph" style:parent-style-name="Heading" style:default-outline-level="1" style:class="text"><style:text-properties fo:font-size="15pt"/></style:style>
    <style:style style:name="Heading_20_2" style:display-name="Heading 2" style:family="paragraph" style:parent-style-name="Heading" style:default-outline-level="2" style:class="text"><style:text-properties fo:font-size="13pt"/></style:style>
    <style:style style:name="Table_20_Contents" style:display-name="Table Contents" style:family="paragraph" style:parent-style-name="Standard" style:class="extra"/>
    <style:style style:name="Table_20_Heading" style:display-name="Table Heading" style:family="paragraph" style:parent-style-name="Table_20_Contents" style:class="extra"><style:text-properties fo:font-weight="bold"/></style:style>
    <style:style style:name="Footnote" style:family="paragraph" style:parent-style-name="Standard" style:class="extra"><style:text-properties fo:font-size="9pt"/></style:style>
    <style:style style:name="Internet_20_link" style:display-name="Internet link" style:family="text"><style:text-properties fo:color="#000080" style:text-underline-style="solid"/></style:style>
    <text:list-style style:name="List_20_1" style:display-name="List 1">
      <text:list-level-style-bullet text:level="1" text:bullet-char="•"><style:list-level-properties text:list-level-position-and-space-mode="label-alignment"><style:list-level-label-alignment text:label-followed-by="listtab" fo:text-indent="-0.4cm" fo:margin-left="0.8cm"/></style:list-level-properties></text:list-level-style-bullet>
      <text:list-level-style-bullet text:level="2" text:bullet-char="◦"><style:list-level-properties text:list-level-position-and-space-mode="label-alignment"><style:list-level-label-alignment text:label-followed-by="listtab" fo:text-indent="-0.4cm" fo:margin-left="1.4cm"/></style:list-level-properties></text:list-level-style-bullet>
    </text:list-style>
    <text:list-style style:name="Numbering_20_123" style:display-name="Numbering 123">
      <text:list-level-style-number text:level="1" style:num-suffix="." style:num-format="1"><style:list-level-properties text:list-level-position-and-space-mode="label-alignment"><style:list-level-label-alignment text:label-followed-by="listtab" fo:text-indent="-0.6cm" fo:margin-left="0.8cm"/></style:list-level-properties></text:list-level-style-number>
    </text:list-style>
    <text:notes-configuration text:note-class="footnote" style:num-format="1" text:start-value="0" text:footnotes-position="page" text:start-numbering-at="document"/>
  </office:styles>
  <office:automatic-styles>
    <style:style style:name="Tab" style:family="table"><style:table-properties style:width="16cm" table:align="margins"/></style:style>
    <style:style style:name="Zelle" style:family="table-cell"><style:table-cell-properties fo:padding="0.1cm" fo:border="0.5pt solid #000000"/></style:style>
    <style:style style:name="fr1" style:family="graphic"><style:graphic-properties style:wrap="none" style:vertical-pos="top" style:horizontal-pos="left" fo:border="none"/></style:style>
    <style:style style:name="fr2" style:family="graphic"><style:graphic-properties style:vertical-pos="top" fo:border="none"/></style:style>
    <style:page-layout style:name="pm1"><style:page-layout-properties fo:page-width="21cm" fo:page-height="29.7cm" fo:margin-top="2cm" fo:margin-bottom="2cm" fo:margin-left="2.5cm" fo:margin-right="2.5cm"/></style:page-layout>
  </office:automatic-styles>
  <office:master-styles><style:master-page style:name="Standard" style:page-layout-name="pm1"/></office:master-styles>
  <office:body>
    <office:text>
      <text:sequence-decls>
        <text:sequence-decl text:display-outline-level="0" text:name="Abbildung"/>
        <text:sequence-decl text:display-outline-level="0" text:name="Tabelle"/>
      </text:sequence-decls>{body}
    </office:text>
  </office:body>
</office:document>
'''


def markdown():
    t2 = ['| ' + ' | '.join(T2_KOPF) + ' |', '|---|---|---:|']
    grid = {}
    for z, s, t, cs, rs in T2_ZELLEN:
        grid[(z, s)] = t
    for z in range(1, 5):
        t2.append('| ' + ' | '.join(grid.get((z, s), '') for s in range(3)) + ' |')
    t1 = ['| ' + ' | '.join(T1_KOPF) + ' |', '|---|---|---:|---|'] + ['| ' + ' | '.join(r) + ' |' for r in T1]
    feuchte = []
    for t, sub in LISTE_FEUCHTE:
        feuchte.append(f'- {t}')
        feuchte += [f'  - {s}' for s in sub]
    return f'''---
titel: {TITEL}
kennung: {ID_REFERENZ}
datum: {DATUM}
autoren:
  - Büro Beispiel Bautechnik
sprache: de-AT
vorlage: bericht
---

# {TITEL}

::: zusammenfassung
{P_ZUSAMMENFASSUNG}
:::

## Befund {{#Befund}}

Begehung am 9. Oktober 2026 mit dem Eigentümer. Die Holzfeuchte wurde an vier Punkten gemessen[^1]; die Werte stehen im Abschnitt [Messwerte](#Messwerte) und in [Tabelle 1](#tabelle-holzfeuchte).

### Feuchteschäden

Feuchteflecken am Sparren Nord und an der Mauerbank, siehe [Abbildung 1](#abbildung-dachstuhl).

![{ALT_FOTO}](bilder/dachstuhl.png "{CAP_FOTO}"){{#abbildung-dachstuhl}}

{chr(10).join(feuchte)}

### Holzschädlinge

{P_HOLZ}

## Messwerte {{#Messwerte}}

Gemessen in 2 cm Tiefe bei 8 °C Lufttemperatur[^2].

{chr(10).join(t1)}

Table: {CAP_T1} {{#tabelle-holzfeuchte}}

![{ALT_DIAGRAMM}](bilder/feuchte-diagramm.png "{CAP_DIAGRAMM}")

## Kosten

{chr(10).join(t2)}

Table: {CAP_T2}

## Empfehlung

{P_EMPFEHLUNG}[Austrian Standards]({URL}).

### Sofortmaßnahmen

1. {LISTE_MASSNAHMEN[0]}
2. {LISTE_MASSNAHMEN[1]}[^3]
3. {LISTE_MASSNAHMEN[2]}

![](bilder/zierlinie.png){{.dekorativ}}

{SCHLUSS}

[^1]: {FN[0]}
[^2]: {FN[1]}
[^3]: {FN[2]}
'''


def erwartet():
    abschnitte = [(2, 'Zusammenfassung'), (2, 'Befund'), (3, 'Feuchteschäden'), (3, 'Holzschädlinge'), (2, 'Messwerte'), (2, 'Kosten'), (2, 'Empfehlung'), (3, 'Sofortmaßnahmen')]
    t1_zellen = [{'zeile': 0, 'spalte': s, 'text': t, 'kopf': True} for s, t in enumerate(T1_KOPF)]
    t1_zellen += [{'zeile': z + 1, 'spalte': s, 'text': t} for z, r in enumerate(T1) for s, t in enumerate(r)]
    t2_zellen = [{'zeile': 0, 'spalte': s, 'text': t, 'kopf': True} for s, t in enumerate(T2_KOPF)]
    t2_zellen += [{'zeile': z, 'spalte': s, 'text': t, **({'colspan': cs} if cs > 1 else {}), **({'rowspan': rs} if rs > 1 else {})} for z, s, t, cs, rs in T2_ZELLEN]
    fliesstext = ' '.join([P_ZUSAMMENFASSUNG, 'Begehung am 9. Oktober 2026 mit dem Eigentümer. Die Holzfeuchte wurde an vier Punkten gemessen; die Werte stehen im Abschnitt Messwerte und in Tabelle 1.',
                           'Feuchteflecken am Sparren Nord und an der Mauerbank, siehe Abbildung 1.', P_HOLZ, 'Gemessen in 2 cm Tiefe bei 8 °C Lufttemperatur.',
                           P_EMPFEHLUNG + 'Austrian Standards.', SCHLUSS] + [t for t, _ in LISTE_FEUCHTE] + [s for _, sub in LISTE_FEUCHTE for s in sub] + LISTE_MASSNAHMEN + FN)
    return {
        'beschreibung': 'Grundwahrheit des Referenzberichts (scripts/s5-quellen.py); gilt für referenz.fodt, referenz.docx, referenz.pdf und referenz.md (Markdown ohne verbundene Zellen).',
        'titel': TITEL,
        'abschnitte': [{'ebene': e, 'text': t} for e, t in abschnitte],
        'tabellen': [
            {'beschriftung': CAP_T1, 'zeilen': 5, 'spalten': 4, 'kopfzeilen': 1, 'zellen': t1_zellen},
            {'beschriftung': CAP_T2, 'zeilen': 5, 'spalten': 3, 'kopfzeilen': 1, 'zellen': t2_zellen},
        ],
        'bilder': [
            {'name': 'dachstuhl', 'alt': ALT_FOTO, 'beschriftung': CAP_FOTO},
            {'name': 'feuchte-diagramm', 'alt': ALT_DIAGRAMM, 'beschriftung': CAP_DIAGRAMM},
            {'name': 'zierlinie', 'dekorativ': True},
        ],
        'fussnoten': [{'text': t} for t in FN],
        'listen': [{'nummeriert': False, 'eintraege': 3, 'unterliste': 2}, {'nummeriert': True, 'eintraege': 3}],
        'verweise': {'intern': [{'text': 'Messwerte', 'ziel': 'Abschnitt Messwerte'}, {'text': 'Tabelle 1', 'ziel': 'Tabelle 1'}, {'text': 'Abbildung 1', 'ziel': 'Abbildung 1'}], 'extern': [URL]},
        'fliesstext': fliesstext,
    }


# --------------------------------------------------------------------------------------------- grenzfall.docx

def grenzfall(foto, zier):
    import docx
    from docx.enum.text import WD_BREAK
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Cm

    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    d = docx.Document()
    cp = d.core_properties
    cp.title = 'Grenzfälle Word-Import'
    cp.author = 'Büro Beispiel'
    cp.language = 'de-AT'
    import datetime
    cp.modified = datetime.datetime(2026, 10, 13, 8, 0, 0)
    cp.created = datetime.datetime(2026, 10, 13, 8, 0, 0)
    cp.identifier = ID_GRENZFALL

    sec = d.sections[0]
    sec.header.paragraphs[0].text = 'Büro Beispiel – Grenzfälle'
    fp = sec.footer.paragraphs[0]
    fp.text = 'Seite '
    fp._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))

    d.add_paragraph('Grenzfälle Word-Import', style='Title')
    # 1. Inhaltsverzeichnis als Feld mit zwischengespeichertem Ergebnis.
    toc = d.add_paragraph()
    toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'))
    toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h </w:instrText></w:r>'))
    toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'))
    toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:t xml:space="preserve">Text und Struktur\t1</w:t></w:r>'))
    toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'))

    h = d.add_heading('Text und Struktur', level=1)
    h._p.insert(1, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="1" w:name="Struktur"/>'))
    h._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="1"/>'))
    p = d.add_paragraph('Absatz mit ')
    p.add_run('Strong-Zeichenvorlage', style='Strong')
    p.add_run(', ')
    p.add_run('direkt fett').bold = True
    p.add_run(', ')
    p.add_run('direkt kursiv').italic = True
    p.add_run(' und hochgestellt m')
    p.add_run('2').font.superscript = True
    p.add_run('.')
    # Fußnote (python-docx kennt keine Fußnoten: Teil von Hand).
    p.add_run(' Mit Fußnote')
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="1"/></w:r>'))
    p.add_run('.')

    d.add_paragraph('Dieser Absatz ist nur direkt fett und groß formatiert.').runs[0].bold = True
    d.paragraphs[-1].runs[0].font.size = docx.shared.Pt(16)
    d.add_paragraph('- getippter Spiegelstrich, keine Liste')
    d.add_paragraph('Echte Aufzählung, Eintrag eins', style='List Bullet')
    d.add_paragraph('Echte Aufzählung, Eintrag zwei', style='List Bullet')
    d.add_paragraph('Unterpunkt', style='List Bullet 2')
    d.add_paragraph('Nummeriert eins', style='List Number')
    d.add_paragraph('Nummeriert zwei', style='List Number')
    d.add_paragraph('Zitat aus dem Gutachten.', style='Quote')

    d.add_heading('Verweise und Felder', level=1)
    p = d.add_paragraph('Querverweis auf ')
    for xml in ['<w:r><w:fldChar w:fldCharType="begin"/></w:r>', '<w:r><w:instrText xml:space="preserve"> REF Struktur \\h </w:instrText></w:r>', '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', '<w:r><w:t>Text und Struktur</w:t></w:r>', '<w:r><w:fldChar w:fldCharType="end"/></w:r>']:
        p._p.append(parse_xml(xml.replace('<w:r>', f'<w:r {nsdecls("w")}>', 1)))
    p.add_run(' auf Seite ')
    for xml in ['<w:r><w:fldChar w:fldCharType="begin"/></w:r>', '<w:r><w:instrText xml:space="preserve"> PAGEREF Struktur \\h </w:instrText></w:r>', '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', '<w:r><w:t>1</w:t></w:r>', '<w:r><w:fldChar w:fldCharType="end"/></w:r>']:
        p._p.append(parse_xml(xml.replace('<w:r>', f'<w:r {nsdecls("w")}>', 1)))
    p.add_run(', Link ')
    rid = d.part.relate_to('https://example.org/normen', RT.HYPERLINK, is_external=True)
    p._p.append(parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}"><w:r><w:t>example.org</w:t></w:r></w:hyperlink>'))
    p.add_run(' und Mail ')
    rid2 = d.part.relate_to('mailto:buero@example.org', RT.HYPERLINK, is_external=True)
    p._p.append(parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid2}"><w:r><w:t>buero@example.org</w:t></w:r></w:hyperlink>'))
    p.add_run('.')
    p = d.add_paragraph('Datum als Feld: ')
    p._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" DATE \\@ &quot;d. MMMM yyyy&quot; "><w:r><w:t>13. Oktober 2026</w:t></w:r></w:fldSimple>'))
    p = d.add_paragraph('Inhaltssteuerelement: ')
    p._p.append(parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="Auftraggeber"/></w:sdtPr><w:sdtContent><w:r><w:t>Maria Beispiel</w:t></w:r></w:sdtContent></w:sdt>'))

    d.add_heading('Überarbeitung', level=1)
    p = d.add_paragraph('Änderungsverfolgung: ')
    p._p.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="11" w:author="Prüfer" w:date="2026-10-12T10:00:00Z"><w:r><w:t xml:space="preserve">eingefügter Text </w:t></w:r></w:ins>'))
    p._p.append(parse_xml(f'<w:del {nsdecls("w")} w:id="12" w:author="Prüfer" w:date="2026-10-12T10:00:00Z"><w:r><w:delText xml:space="preserve">gelöschter Text </w:delText></w:r></w:del>'))
    p.add_run('bleibt.')
    p = d.add_paragraph('Kommentierter Satz.')
    d.add_comment(p.runs, text='Bitte Messwert prüfen.', author='Prüfer', initials='PR')

    d.add_heading('Objekte', level=1)
    p = d.add_paragraph('Formel: ')
    p._p.append(parse_xml(f'<m:oMath {nsdecls("m", "w")}><m:r><m:t>U</m:t></m:r><m:r><m:t>=</m:t></m:r><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>R</m:t></m:r></m:den></m:f></m:oMath>'))
    p = d.add_paragraph('Textfeld folgt:')
    p._p.append(parse_xml(
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape id="Textfeld1" style="width:200pt;height:40pt" type="#_x0000_t202">'
        f'<v:textbox><w:txbxContent><w:p><w:r><w:t>Inhalt des Textfelds</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'))
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    pic = d.add_picture(io.BytesIO(foto), width=Cm(8))
    d.paragraphs[-1].style = d.styles['Normal']
    # Bild mit Alternativtext.
    ip = d.inline_shapes[-1]._inline
    ip.docPr.set('descr', ALT_FOTO)
    cap = d.add_paragraph('Abbildung ', style='Caption')
    cap._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" SEQ Abbildung \\* ARABIC "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    cap.add_run(': Bild mit Alternativtext')
    # Bild ohne Alternativtext.
    d.add_picture(io.BytesIO(foto), width=Cm(4))
    # Als dekorativ markiertes Bild (Office 2019: a:extLst mit adec:decorative).
    d.add_picture(io.BytesIO(zier), width=Cm(10))
    d.inline_shapes[-1]._inline.docPr.append(parse_xml(
        '<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}">'
        '<adec:decorative xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="1"/></a:ext></a:extLst>'))
    # BMP-Bild (kein Kernmedientyp).
    bmp = io.BytesIO()
    Image.open(io.BytesIO(foto)).resize((60, 40)).save(bmp, 'BMP')
    d.add_picture(io.BytesIO(bmp.getvalue()), width=Cm(2))
    d.inline_shapes[-1]._inline.docPr.set('descr', 'Kleines Bild im BMP-Format')

    d.add_heading('Tabellen', level=1)
    t = d.add_table(rows=4, cols=3)
    t.style = 'Table Grid'
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    for i, v in enumerate(['Gewerk', 'Leistung', 'Betrag']):
        t.cell(0, i).text = v
    t.cell(1, 0).text = 'Spengler'
    t.cell(1, 1).text = 'Rinne'
    t.cell(1, 2).text = '1.250,00'
    t.cell(2, 1).text = 'Kamin'
    t.cell(2, 2).text = '640,00'
    a = t.cell(1, 0).merge(t.cell(2, 0))
    a.text = 'Spengler'
    s = t.cell(3, 0).merge(t.cell(3, 1))
    s.text = 'Summe netto'
    t.cell(3, 2).text = '1.890,00'
    cap = d.add_paragraph('Tabelle ', style='Caption')
    cap._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" SEQ Tabelle \\* ARABIC "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    cap.add_run(': Tabelle mit verbundenen Zellen')
    t2 = d.add_table(rows=1, cols=2)
    t2.style = 'Table Grid'
    t2.cell(0, 0).text = 'Außen'
    inner = t2.cell(0, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = 'innen 1'
    inner.cell(0, 1).text = 'innen 2'
    d.add_paragraph('Ende der Grenzfälle.')

    buf = io.BytesIO()
    d.save(buf)
    # Fußnotenteil ergänzen (python-docx 1.2 schreibt keinen).
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    footnotes = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes xmlns:w="{W}">'
                 '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                 '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
                 '<w:footnote w:id="1"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
                 '<w:r><w:t xml:space="preserve"> Text der Fußnote im Word-Dokument.</w:t></w:r></w:p></w:footnote></w:footnotes>')
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            if item.filename == 'word/_rels/document.xml.rels':
                data = data.replace(b'</Relationships>', b'<Relationship Id="rIdFn1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
            zi = zipfile.ZipInfo(item.filename, date_time=(2026, 10, 13, 8, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, data)
        zi = zipfile.ZipInfo('word/footnotes.xml', date_time=(2026, 10, 13, 8, 0, 0))
        zi.compress_type = zipfile.ZIP_DEFLATED
        zout.writestr(zi, footnotes.encode('utf-8'))
    return out.getvalue()


GRENZFALL_ERWARTET = [
    {'merkmal': 'Inhaltsverzeichnisfeld (TOC)', 'erwartet': 'entfällt, das Navigationsdokument ersetzt es; gemeldet', 'verlust': False},
    {'merkmal': 'Überschrift mit Textmarke „Struktur“', 'erwartet': 'h2 id="Struktur"', 'verlust': False},
    {'merkmal': 'Zeichenvorlage „Strong“', 'erwartet': 'strong', 'verlust': False},
    {'merkmal': 'direkt fett / direkt kursiv', 'erwartet': 'span mit CSS-Klasse (keine Struktur aus dem Aussehen, HP-MIN-012); gemeldet', 'verlust': 'Bedeutung'},
    {'merkmal': 'hochgestellt (w:vertAlign)', 'erwartet': 'sup', 'verlust': False},
    {'merkmal': 'Fußnote', 'erwartet': 'noteref + aside', 'verlust': False},
    {'merkmal': 'direkt fett und groß formatierter Absatz', 'erwartet': 'p, keine Überschrift', 'verlust': False},
    {'merkmal': 'getippter Spiegelstrich', 'erwartet': 'p mit Text „- …“, keine Liste', 'verlust': False},
    {'merkmal': 'Aufzählung über Vorlage „List Bullet“ mit Unterpunkt „List Bullet 2“', 'erwartet': 'ul mit verschachtelter ul', 'verlust': False},
    {'merkmal': 'Nummerierung über Vorlage „List Number“', 'erwartet': 'ol', 'verlust': False},
    {'merkmal': 'Zitatvorlage „Quote“', 'erwartet': 'blockquote', 'verlust': False},
    {'merkmal': 'REF-Feld auf Textmarke', 'erwartet': 'a href="#Struktur"', 'verlust': False},
    {'merkmal': 'PAGEREF-Feld', 'erwartet': 'a href="#Struktur" mit Text des Ziels statt Seitenzahl; gemeldet', 'verlust': 'Seitenzahl'},
    {'merkmal': 'Hyperlink https', 'erwartet': 'a href', 'verlust': False},
    {'merkmal': 'Hyperlink mailto', 'erwartet': 'Text, gemeldet (HPR-MIN-010 erlaubt nur http/https)', 'verlust': 'Verknüpfung'},
    {'merkmal': 'DATE-Feld', 'erwartet': 'letztes Ergebnis als Text, gemeldet', 'verlust': 'Feldfunktion'},
    {'merkmal': 'Inhaltssteuerelement', 'erwartet': 'Inhalt als Text, gemeldet', 'verlust': 'Steuerelement'},
    {'merkmal': 'Änderungsverfolgung (Einfügung, Löschung)', 'erwartet': 'Einfügung bleibt, Löschung entfällt (angenommener Stand); gemeldet', 'verlust': 'Verlauf'},
    {'merkmal': 'Kommentar', 'erwartet': 'entfällt, gemeldet', 'verlust': True},
    {'merkmal': 'Formel (OMML)', 'erwartet': 'linearer Text, gemeldet', 'verlust': 'Formelstruktur'},
    {'merkmal': 'Textfeld (VML)', 'erwartet': 'Text als Absatz, gemeldet', 'verlust': 'Position'},
    {'merkmal': 'Seitenumbruch', 'erwartet': 'entfällt', 'verlust': False},
    {'merkmal': 'Bild mit Alternativtext und Beschriftung (SEQ)', 'erwartet': 'figure mit alt und figcaption', 'verlust': False},
    {'merkmal': 'Bild ohne Alternativtext', 'erwartet': 'figure, alt aus dem Namen, gemeldet', 'verlust': 'Alternativtext'},
    {'merkmal': 'als dekorativ markiertes Bild (adec:decorative)', 'erwartet': 'figure mit alt=""', 'verlust': False},
    {'merkmal': 'BMP-Bild', 'erwartet': 'Abbruch mit Nennung; mit --bild-auslassen Platzhalter und Meldung', 'verlust': True},
    {'merkmal': 'Tabelle mit Kopfzeile (tblHeader), gridSpan, vMerge, Beschriftung', 'erwartet': 'thead, colspan, rowspan, caption', 'verlust': False},
    {'merkmal': 'verschachtelte Tabelle', 'erwartet': 'table in td', 'verlust': False},
    {'merkmal': 'Kopf- und Fußzeile', 'erwartet': 'Kopfzeilentext einmal als Absatz am Ende (Teil 5, Kapitel 4.7), Seitenzahlzeile entfällt; gemeldet', 'verlust': 'Wiederholung je Seite'},
]


def main():
    os.makedirs(os.path.join(OUT, 'bilder'), exist_ok=True)
    foto, diagramm, zier = bilder()
    for name, data in [('dachstuhl.png', foto), ('feuchte-diagramm.png', diagramm), ('zierlinie.png', zier)]:
        with open(os.path.join(OUT, 'bilder', name), 'wb') as f:
            f.write(data)
    with open(os.path.join(OUT, 'referenz.fodt'), 'w', encoding='utf-8') as f:
        f.write(fodt(foto, diagramm, zier))
    with open(os.path.join(OUT, 'referenz.md'), 'w', encoding='utf-8') as f:
        f.write(markdown())
    with open(os.path.join(OUT, 'erwartet.json'), 'w', encoding='utf-8') as f:
        json.dump(erwartet(), f, ensure_ascii=False, indent=2)
        f.write('\n')
    with open(os.path.join(OUT, 'grenzfall.docx'), 'wb') as f:
        f.write(grenzfall(foto, zier))
    with open(os.path.join(OUT, 'grenzfall-erwartet.json'), 'w', encoding='utf-8') as f:
        json.dump({'beschreibung': 'Merkmale von grenzfall.docx und ihre erwartete Behandlung (scripts/s5-quellen.py)', 'merkmale': GRENZFALL_ERWARTET}, f, ensure_ascii=False, indent=2)
        f.write('\n')
    for name in sorted(os.listdir(OUT)):
        p = os.path.join(OUT, name)
        if os.path.isfile(p):
            print(f'{name}  {os.path.getsize(p)} Bytes')


if __name__ == '__main__':
    main()
